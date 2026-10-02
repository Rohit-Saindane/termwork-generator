"""
TermWorkForge — Term-Work Filler  (v6)

What changed in v6
------------------
1. Robust index parsing
   * Word spacing is rebuilt from real character gaps instead of pdfplumber's
     fixed 3 pt tolerance, so index PDFs that carry NO space glyphs (e.g.
     "Writeaprogramtoimplement...") are read correctly.
   * The Title / Sr.No. columns are found from the header row (with sensible
     fall-backs), so column order and extra columns no longer matter.
   * Hyphenated line breaks ("Tic-tac-\\ntoe") are re-joined properly.
   * If the index has no ruled table at all, a text-line parser takes over.
2. Alignment-proof PDF writing
   * All text is sanitised to what the Times-Roman font can really draw.
   * The experiment title is wrapped over the two ruled lines using true font
     metrics and auto-shrinks (12 pt -> 7.5 pt) so it always fits inside the
     lines; only if it is absurdly long does it get an ellipsis.
   * Every other field auto-shrinks the same way instead of being chopped.
   * The overlay uses the template's real page size / origin.
3. Downloads are branded:  TermWorkForge_Practicals_2-5.pdf
"""
import io, json, os, re, secrets, statistics, unicodedata

import pdfplumber
from flask import Flask, jsonify, render_template, request, send_file
from pypdf import PdfReader, PdfWriter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from PIL import Image

app = Flask(__name__)
app.secret_key = secrets.token_hex(16)
app.config["MAX_CONTENT_LENGTH"] = 32 * 1024 * 1024  # 32 MB

# ── Branding ───────────────────────────────────────────────────────────────────
BRAND_NAME = "TermWorkForge"          # used in the downloaded file name

# ── Page geometry ──────────────────────────────────────────────────────────────
PAGE_W, PAGE_H = 594.96, 842.52
FONT, SIZE = "Times-Roman", 12
MIN_SIZE = 7.5                        # smallest size text may shrink to

# ── Supported file types ───────────────────────────────────────────────────────
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".webp", ".tiff", ".tif"}
DOCX_EXTS = {".docx", ".doc"}
PDF_EXT = {".pdf"}
ALL_TERM_WORK_EXTS = PDF_EXT | IMAGE_EXTS | DOCX_EXTS
ALL_INDEX_EXTS = PDF_EXT  # Index must be PDF — OCR on Render free tier not possible

# ── Field injection coordinates (x, y_reportlab, max_x) ───────────────────────
FIELDS = {
    "term":           (502, 609.2, 527),
    "subject":        (111, 585.4, 527),
    "pen":            (97,  562.0, 385),
    "semester":       (479, 562.0, 527),
    "student_name":   (157, 536.9, 526),
    "class_name":     (101, 516.6, 386),
    "batch":          (463, 516.6, 525),
    "exp_z1":         (283, 494.9, 523),
    "exp_z2":         (73,  469.9, 526),
    "checked_by":     (136, 263.2, 524),
    "practical_no":   (407, 406.6, 531),
    "cond_dd":        (415, 383.7, 445),
    "cond_mm":        (450, 383.7, 480),
    "cond_yyyy":      (485, 383.7, 517),
    "sub_dd":         (413, 361.2, 444),
    "sub_mm":         (448, 360.2, 480),
    "sub_yyyy":       (484, 360.2, 515),
    "act_dd":         (414, 337.7, 447),
    "act_mm":         (451, 337.7, 484),
    "act_yyyy":       (488, 337.7, 518),
    "sign_date_dd":   (290, 97.5,  322),
    "sign_date_mm":   (327, 97.5,  357),
    "sign_date_yyyy": (363, 97.9,  392),
    "marks":          (478, 102.1, 525),
}


# ─────────────────────────────────────────────────────────────────────────────
# FILE CONVERSION HELPERS
# ─────────────────────────────────────────────────────────────────────────────
def _ext(filename: str) -> str:
    return os.path.splitext(filename.lower())[1] if filename else ""


def image_bytes_to_pdf(img_bytes: bytes) -> bytes:
    """
    Convert any PIL-readable image (JPG/PNG/BMP/WEBP/TIFF) to a single-page PDF.
    The image is scaled to fit A4 with margins, centred, white background.
    """
    img = Image.open(io.BytesIO(img_bytes)).convert("RGB")
    iw, ih = img.size
    W, H = A4
    margin = 28
    avail_w = W - 2 * margin
    avail_h = H - 2 * margin
    scale = min(avail_w / iw, avail_h / ih)
    dw, dh = iw * scale, ih * scale
    x = margin + (avail_w - dw) / 2
    y = margin + (avail_h - dh) / 2

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFillColorRGB(1, 1, 1)
    c.rect(0, 0, W, H, fill=1, stroke=0)
    img_buf = io.BytesIO()
    img.save(img_buf, "PNG")
    img_buf.seek(0)
    c.drawImage(ImageReader(img_buf), x, y, dw, dh)
    c.save()
    buf.seek(0)
    return buf.read()


def docx_to_pdf(docx_bytes: bytes) -> bytes:
    """
    Convert a DOCX file to PDF by extracting all text and reflowing into
    a plain-text PDF. Preserves structure (headings, paragraphs, tables).

    Note: this is a text extraction approach — full layout fidelity isn't
    possible without LibreOffice, which isn't available on Render free tier.
    """
    import docx as python_docx
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors

    doc = python_docx.Document(io.BytesIO(docx_bytes))
    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=50, rightMargin=50,
                            topMargin=50, bottomMargin=50)
    styles = getSampleStyleSheet()
    story = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            story.append(Spacer(1, 6))
            continue
        style_name = para.style.name or ""
        if "Heading 1" in style_name:
            story.append(Paragraph(text, styles["Heading1"]))
        elif "Heading 2" in style_name:
            story.append(Paragraph(text, styles["Heading2"]))
        else:
            story.append(Paragraph(text, styles["Normal"]))
        story.append(Spacer(1, 4))

    for table in doc.tables:
        data = [[cell.text for cell in row.cells] for row in table.rows]
        if data:
            t = Table(data)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("PADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(t)
            story.append(Spacer(1, 8))

    if not story:
        story.append(Paragraph("(Empty document)", styles["Normal"]))

    pdf.build(story)
    buf.seek(0)
    return buf.read()


def to_pdf_bytes(file_bytes: bytes, filename: str) -> bytes:
    """Route any supported file format to PDF bytes."""
    ext = _ext(filename)
    if ext in PDF_EXT:
        return file_bytes
    if ext in IMAGE_EXTS:
        return image_bytes_to_pdf(file_bytes)
    if ext in DOCX_EXTS:
        return docx_to_pdf(file_bytes)
    raise ValueError(f"Unsupported file type: {ext}")


def validate_ext(filename: str, allowed: set, field_name: str):
    """Return (ok, error_message)."""
    ext = _ext(filename)
    if ext not in allowed:
        nice = ", ".join(sorted(allowed))
        return False, f"{field_name} must be one of: {nice}. Got '{ext}'."
    return True, ""


# ─────────────────────────────────────────────────────────────────────────────
# TEXT SAFETY  (make sure whatever we draw exists in the Times-Roman font)
# ─────────────────────────────────────────────────────────────────────────────
_CHAR_MAP = {
    "\u2010": "-", "\u2011": "-", "\u2012": "-", "\u2013": "-", "\u2014": "-",
    "\u2015": "-", "\u2212": "-", "\u00ad": "",
    "\u2018": "'", "\u2019": "'", "\u201a": ",", "\u201b": "'",
    "\u201c": '"', "\u201d": '"', "\u201e": '"',
    "\u00a0": " ", "\u2009": " ", "\u200a": " ", "\u202f": " ",
    "\u2022": "-", "\u25cf": "-", "\u00b7": "-",
}


def sanitize(text) -> str:
    """
    Normalise any string so it can be drawn safely in Times-Roman:
    ligatures (ﬁ -> fi), fancy dashes/quotes, control chars, non-Latin glyphs
    and runs of whitespace are all handled.
    """
    if text is None:
        return ""
    text = unicodedata.normalize("NFKC", str(text))
    out = []
    for ch in text:
        ch = _CHAR_MAP.get(ch, ch)
        for c in ch:
            if c in "\n\r\t\v\f":
                out.append(" ")
                continue
            if unicodedata.category(c) in ("Cc", "Cf", "Cs", "Co", "Cn"):
                continue
            try:
                c.encode("cp1252")
            except UnicodeEncodeError:
                continue
            out.append(c)
    return re.sub(r"\s+", " ", "".join(out)).strip()


# ─────────────────────────────────────────────────────────────────────────────
# TEXT FITTING  (never overflow, never overlap — shrink instead)
# ─────────────────────────────────────────────────────────────────────────────
def _w(text, size=SIZE):
    return pdfmetrics.stringWidth(text, FONT, size)


def _size_steps(start=SIZE, stop=MIN_SIZE, step=0.5):
    s = float(start)
    while s >= stop - 1e-9:
        yield round(s, 2)
        s -= step


def _ellipsize(text, max_w, size):
    """Cut text so that text + '...' fits inside max_w."""
    if _w(text, size) <= max_w:
        return text
    ell = "..."
    while text and _w(text + ell, size) > max_w:
        text = text[:-1]
    text = text.rstrip()
    return (text + ell) if text else ""


def _hard_split(word, max_w, size):
    """Break one over-long word into chunks that each fit max_w."""
    chunks, cur = [], ""
    for ch in word:
        if cur and _w(cur + ch, size) > max_w:
            chunks.append(cur)
            cur = ch
        else:
            cur += ch
    if cur:
        chunks.append(cur)
    return chunks


def fit_single_line(text, max_w, start=SIZE):
    """Return (text, size): largest size <= start where text fits; else ellipsis at MIN_SIZE."""
    text = sanitize(text)
    for s in _size_steps(start):
        if _w(text, s) <= max_w:
            return text, s
    return _ellipsize(text, max_w, MIN_SIZE), MIN_SIZE


def wrap_experiment(title):
    """
    Lay the experiment title over the two ruled lines of the form.
    Returns (line1, line2, font_size).  Both lines share one font size so the
    result always looks uniform.
    """
    title = sanitize(title)
    if not title:
        return "", "", SIZE

    x1, _, m1 = FIELDS["exp_z1"]
    x2, _, m2 = FIELDS["exp_z2"]
    w1, w2 = m1 - x1, m2 - x2

    def layout(size):
        words = title.split(" ")
        # A single word wider than a whole line would break everything: split it.
        flat = []
        for wd in words:
            flat.extend(_hard_split(wd, w2, size) if _w(wd, size) > w2 else [wd])
        l1 = ""
        i = 0
        while i < len(flat):
            cand = (l1 + " " + flat[i]).strip()
            if _w(cand, size) <= w1:
                l1 = cand
                i += 1
            else:
                break
        rest = " ".join(flat[i:])
        return l1, rest, _w(rest, size) <= w2

    fallback = None
    for s in _size_steps():
        l1, l2, ok = layout(s)
        if ok:
            if l1:                       # nice: starts on the first line
                return l1, l2, s
            if fallback is None:         # fits, but the first line stayed empty
                fallback = (l1, l2, s)
    if fallback:
        return fallback

    # Truly enormous title: last resort, smallest size + ellipsis on line 2
    l1, l2, _ = layout(MIN_SIZE)
    return l1, _ellipsize(l2, w2, MIN_SIZE), MIN_SIZE


# ─────────────────────────────────────────────────────────────────────────────
# INDEX PARSING  (robust against odd PDFs)
# ─────────────────────────────────────────────────────────────────────────────
def _gap_threshold(gaps):
    """
    Gap (in em) that counts as a word space. In normal fonts the gap inside a
    word is ~0 (kerning is < 0.1 em) while a real space is ~0.25 em, so 0.15 em
    cleanly separates the two — and, unlike a fixed 3 pt tolerance, it scales
    with the font size.
    """
    return 0.15


def _dedupe_chars(chars):
    """Some PDFs fake bold by printing every glyph twice — drop the duplicates."""
    seen, out = set(), []
    for c in chars:
        key = (c["text"], round(c["x0"] / 0.5), round(c["top"] / 0.5))
        if key in seen:
            continue
        seen.add(key)
        out.append(c)
    return out


def _group_lines(chars):
    """Cluster characters into visual lines (top -> bottom)."""
    chars = sorted(_dedupe_chars(chars), key=lambda c: (c["top"], c["x0"]))
    lines = []
    for ch in chars:
        h = ch.get("size") or (ch["bottom"] - ch["top"]) or 10
        mid = (ch["top"] + ch["bottom"]) / 2
        if lines and abs(mid - lines[-1]["mid"]) <= 0.5 * h:
            ln = lines[-1]
            ln["chars"].append(ch)
            ln["mid"] += (mid - ln["mid"]) / len(ln["chars"])
        else:
            lines.append({"mid": mid, "chars": [ch], "size": h})
    return lines


def _line_text(chars):
    """Rebuild one line of text, inserting spaces from real glyph gaps."""
    chars = sorted(chars, key=lambda c: c["x0"])
    if not chars:
        return ""
    sizes = [c.get("size") or 10 for c in chars]
    size = statistics.median(sizes) or 10
    gaps = [(b["x0"] - a["x1"]) / size for a, b in zip(chars, chars[1:])]
    thr = _gap_threshold(gaps)
    out = []
    for i, ch in enumerate(chars):
        t = ch["text"]
        if t.isspace():
            out.append(" ")
            continue
        if i > 0 and gaps[i - 1] > thr:
            out.append(" ")
        out.append(t)
    return re.sub(r"\s+", " ", "".join(out)).strip()


def _join_lines(parts):
    """Join wrapped lines; 'Tic-tac-' + 'toe' becomes 'Tic-tac-toe'."""
    text = ""
    for ln in parts:
        if not ln:
            continue
        if not text:
            text = ln
        elif text.endswith("-") and not text.endswith(" -") and ln[:1].isalnum():
            text += ln
        else:
            text += " " + ln
    return re.sub(r"\s+", " ", text).strip()


def _chars_to_text(chars):
    return _join_lines([_line_text(l["chars"]) for l in _group_lines(chars)])


def _chars_in_bbox(chars, bbox):
    x0, top, x1, bottom = bbox
    res = []
    for c in chars:
        cx = (c["x0"] + c["x1"]) / 2
        cy = (c["top"] + c["bottom"]) / 2
        if x0 <= cx <= x1 and top <= cy <= bottom:
            res.append(c)
    return res


# ── table → practicals ────────────────────────────────────────────────────────
_TITLE_KEYS = ["title", "experiment", "aim", "practical", "program",
               "problem", "description", "topic", "name"]
_SR_RE = re.compile(r"^\s*(\d{1,3})\s*[.)]?\s*$")
_DATE_RE = re.compile(r"^\d{1,4}\s*[/\-.]\s*\d{1,2}\s*[/\-.]\s*\d{1,4}$")


def _has_letters(s):
    return bool(re.search(r"[A-Za-z]{2,}", s or ""))


_SR_HEAD_RE = re.compile(
    r"\b(sr|sl|s)\b\.?\s*(no|num)?|serial|^\s*(exp|prac\w*)?\.?\s*no\.?\s*$|^\s*#\s*$|\bexp\.?\s*no",
    re.IGNORECASE,
)


def _find_columns(grid):
    """Locate the Sr.No. and Title columns (header hints first, then by looking at the data)."""
    ncols = max((len(r) for r in grid), default=0)

    # --- Sr.No. column: the column that mostly holds small whole numbers ------
    def num_hits(ci):
        return sum(1 for r in grid if ci < len(r) and _SR_RE.match((r[ci] or "").strip()))

    sr_col = None
    for row in grid[:4]:
        for ci, cell in enumerate(row):
            if cell and _SR_HEAD_RE.search(cell.strip()) and num_hits(ci) >= 2:
                sr_col = ci
                break
        if sr_col is not None:
            break
    if sr_col is None and ncols:
        best = max(range(ncols), key=num_hits)
        if num_hits(best) >= 2:
            sr_col = best

    # --- Title column: header keyword, else the column with the most text -----
    title_col = None
    for row in grid[:4]:
        lows = [(c or "").lower() for c in row]
        for key in _TITLE_KEYS:
            for ci, low in enumerate(lows):
                if ci != sr_col and key in low:
                    title_col = ci
                    break
            if title_col is not None:
                break
        if title_col is not None:
            break
    if title_col is None and ncols:
        def text_len(ci):
            return sum(len(r[ci] or "") for r in grid
                       if ci < len(r) and ci != sr_col and _has_letters(r[ci]))
        cand = max(range(ncols), key=text_len)
        if text_len(cand) > 0 and cand != sr_col:
            title_col = cand
    return sr_col, title_col


def _parse_grid(grid):
    """grid: list of rows, each a list of cell strings → {sr_no: title}."""
    if not grid:
        return {}
    sr_col, title_col = _find_columns(grid)
    if sr_col is None:
        return {}

    found, last = {}, None
    for row in grid:
        if not row or sr_col >= len(row):
            continue
        sr_cell = (row[sr_col] or "").strip()
        m = _SR_RE.match(sr_cell)

        if m:
            title = ""
            if title_col is not None and title_col < len(row):
                title = row[title_col] or ""
            if not _has_letters(title):
                # fall back: the longest "wordy" cell of the row
                cands = [(c or "") for i, c in enumerate(row)
                         if i != sr_col and _has_letters(c) and not _DATE_RE.match((c or "").strip())]
                title = max(cands, key=len) if cands else ""
            title = sanitize(title)
            if title:
                n = int(m.group(1))
                if n not in found:
                    found[n] = title
                    last = n
        elif last is not None and not sr_cell and title_col is not None and title_col < len(row):
            # wrapped title continuing on the next row (every other cell empty)
            others = [c for i, c in enumerate(row) if i != title_col and (c or "").strip()]
            cont = row[title_col] or ""
            if not others and _has_letters(cont):
                found[last] = sanitize(found[last] + " " + cont)
    return found


def _extract_from_tables(page):
    found = {}
    try:
        tables = page.find_tables()
    except Exception:
        return found
    chars = page.chars
    for tb in tables:
        grid = []
        for r in tb.rows:
            row = []
            for bbox in r.cells:
                row.append("" if bbox is None else _chars_to_text(_chars_in_bbox(chars, bbox)))
            grid.append(row)
        for k, v in _parse_grid(grid).items():
            found.setdefault(k, v)
    return found


# ── text-only fallback (index without ruled table) ────────────────────────────
_LINE_RE = re.compile(
    r"^\s*(?:practical|experiment|exp\.?|prac\.?|program)?\s*(\d{1,3})\s*[.)\-:]?\s+(\S.*)$",
    re.IGNORECASE,
)


def _extract_from_text(page):
    found, last, last_mid = {}, None, None
    for ln in _group_lines(page.chars):
        text = _line_text(ln["chars"])
        if not text:
            continue
        m = _LINE_RE.match(text)
        if m and _has_letters(m.group(2)):
            n = int(m.group(1))
            if n not in found:
                found[n] = sanitize(m.group(2))
                last, last_mid = n, ln["mid"]
                continue
        if last is not None and last_mid is not None:
            close = (ln["mid"] - last_mid) < 2.0 * ln["size"]
            prev = found[last]
            if close and not _DATE_RE.match(text):
                found[last] = sanitize(_join_lines([prev, text]))
                last_mid = ln["mid"]
                continue
        last = None
    return found


def extract_practicals(file_bytes: bytes, filename: str = "index.pdf") -> dict:
    """
    Extract {sr_no: title} from a PDF index.
    1) ruled tables (header-aware)  2) plain text lines as a fall-back.
    Only PDF is supported — image OCR needs the tesseract binary which isn't
    available on Render free tier.
    """
    practicals = {}
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            for k, v in _extract_from_tables(page).items():
                practicals.setdefault(k, v)
        if len(practicals) < 1:
            for page in pdf.pages:
                for k, v in _extract_from_text(page).items():
                    practicals.setdefault(k, v)
        if len(practicals) < 2 and len(practicals) >= 1:
            # a lone hit from the table pass could be noise — try text mode too
            extra = {}
            for page in pdf.pages:
                extra.update({k: v for k, v in _extract_from_text(page).items()})
            if len(extra) > len(practicals):
                practicals = extra
    return dict(sorted(practicals.items()))


_SUBJECT_RE = re.compile(r"^\s*S\s*U\s*B\s*J\s*E\s*C\s*T\s*[:\-]\s*(.*)$", re.IGNORECASE)
_HEADER_WORDS = re.compile(r"^(sr\b|s\.?\s*no|no\.?\b|date\b|title\b|pen\b|name\b|rubrics\b)", re.IGNORECASE)


def extract_subject_from_index(pdf_bytes: bytes) -> str:
    """
    Extract the Subject name from the practical index PDF.
    Returns "" when the Subject line is empty (never grabs the table header).
    """
    try:
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            for page in pdf.pages:
                for ln in _group_lines(page.chars):
                    text = _line_text(ln["chars"])
                    m = _SUBJECT_RE.match(text)
                    if not m:
                        continue
                    subject = m.group(1).strip()
                    if not subject or _HEADER_WORDS.match(subject):
                        return ""
                    # Fix OCR letter-splits such as "O perating" → "Operating"
                    subject = re.sub(r"\b([B-HJ-Z])\s(?=[a-z]{3,}\b)", r"\1", subject)
                    return sanitize(subject)
    except Exception:
        pass
    return ""


# ─────────────────────────────────────────────────────────────────────────────
# PDF BUILDING
# ─────────────────────────────────────────────────────────────────────────────
def build_pdf(pdf_b_bytes: bytes, title: str, form: dict) -> bytes:
    reader_b = PdfReader(io.BytesIO(pdf_b_bytes))
    first = reader_b.pages[0]

    # Use the template's real page size / origin so the overlay lines up exactly
    try:
        left, bottom = float(first.mediabox.left), float(first.mediabox.bottom)
        pw, ph = float(first.mediabox.width), float(first.mediabox.height)
    except Exception:
        left, bottom, pw, ph = 0.0, 0.0, PAGE_W, PAGE_H

    l1, l2, exp_size = wrap_experiment(title)

    packet = io.BytesIO()
    c = canvas.Canvas(packet, pagesize=(pw + left, ph + bottom))
    if left or bottom:
        c.translate(left, bottom)

    def put(key, text):
        text = sanitize(text)
        if not text:
            return
        x, y, mx = FIELDS[key]
        txt, size = fit_single_line(text, mx - x)
        c.setFont(FONT, size)
        c.drawString(x, y, txt)

    put("term", form.get("term", ""))
    put("subject", form.get("subject", ""))
    put("pen", form.get("pen", ""))
    put("semester", form.get("semester", ""))
    put("student_name", form.get("student_name", ""))
    put("class_name", form.get("class_name", ""))
    put("batch", form.get("batch", ""))
    put("checked_by", form.get("checked_by", ""))
    put("practical_no", form.get("practical_no", ""))

    def put_date(prefix, date_str):
        if not date_str:
            return
        parts = re.split(r"[/\-\.]", str(date_str).strip())
        if len(parts) == 3:
            put(f"{prefix}_dd", parts[0].strip().zfill(2))
            put(f"{prefix}_mm", parts[1].strip().zfill(2))
            put(f"{prefix}_yyyy", parts[2].strip())

    put_date("cond", form.get("conducted_on", ""))
    put_date("sub", form.get("date_of_sub", ""))
    put_date("act", form.get("actual_sub", ""))
    put_date("sign_date", form.get("sign_date", ""))

    put("marks", form.get("marks", ""))

    # Experiment title — both lines share one font size
    c.setFont(FONT, exp_size)
    if l1:
        c.drawString(FIELDS["exp_z1"][0], FIELDS["exp_z1"][1], l1)
    if l2:
        c.drawString(FIELDS["exp_z2"][0], FIELDS["exp_z2"][1], l2)

    c.save()
    packet.seek(0)
    overlay = PdfReader(packet).pages[0]

    writer = PdfWriter()
    first.merge_page(overlay)
    writer.add_page(first)
    for i in range(1, len(reader_b.pages)):
        writer.add_page(reader_b.pages[i])

    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_download_name(sr_numbers) -> str:
    """[2,3,4,5] → TermWorkForge_Practicals_2-5.pdf ; [1,3,4,5,9] → …_1_3-5_9.pdf"""
    nums = sorted({int(n) for n in sr_numbers})
    if not nums:
        return f"{BRAND_NAME}_Practicals.pdf"
    groups, start, prev = [], nums[0], nums[0]
    for n in nums[1:]:
        if n == prev + 1:
            prev = n
            continue
        groups.append((start, prev))
        start = prev = n
    groups.append((start, prev))
    parts = []
    for a, b in groups:
        if a == b:
            parts.append(str(a))
        elif b == a + 1:
            parts.extend([str(a), str(b)])
        else:
            parts.append(f"{a}-{b}")
    label = "Practical" if len(nums) == 1 else "Practicals"
    return f"{BRAND_NAME}_{label}_{'_'.join(parts)}.pdf"


# ─────────────────────────────────────────────────────────────────────────────
# ROUTES
# ─────────────────────────────────────────────────────────────────────────────
@app.route("/")
def index():
    return render_template("index.html")


@app.route("/form")
def form_page():
    return render_template("form.html")


@app.route("/health")
def health():
    resp = jsonify({"status": "ok", "service": "termworkforge"})
    resp.headers["Access-Control-Allow-Origin"] = "*"
    resp.headers["Access-Control-Allow-Methods"] = "GET, HEAD, OPTIONS"
    resp.headers["Cache-Control"] = "no-store"
    return resp


@app.route("/template/<name>")
def get_template(name):
    safe = re.sub(r"[^a-zA-Z0-9_\-]", "", name)
    path = os.path.join(app.root_path, "static", "templates", safe + ".pdf")
    if not os.path.exists(path):
        return jsonify({"error": "Template not found"}), 404
    return send_file(path, mimetype="application/pdf")


@app.route("/api/parse-index", methods=["POST"])
def parse_index():
    if "pdf_a" not in request.files:
        return jsonify({"error": "No index file uploaded"}), 400
    f = request.files["pdf_a"]
    filename = f.filename or "index.pdf"
    ok, err = validate_ext(filename, ALL_INDEX_EXTS, "Practical Index")
    if not ok:
        return jsonify({"error": err}), 400
    try:
        file_bytes = f.read()
        data = extract_practicals(file_bytes, filename)
        if not data:
            return jsonify({"error": "No practicals found. Make sure the index has a table with SR. No. and Title columns."}), 422
        # Also extract subject — returned alongside practicals so form.js can auto-fill it
        subject = extract_subject_from_index(file_bytes)
        return jsonify({"practicals": data, "subject": subject})
    except Exception as e:
        msg = str(e)
        if "Root object" in msg or "EOF" in msg:
            msg = "This file doesn't look like a valid PDF. Please check and try again."
        return jsonify({"error": msg}), 500


@app.route("/api/generate", methods=["POST"])
def generate():
    if "pdf_b" not in request.files:
        return jsonify({"error": "No Term Work file uploaded"}), 400
    f = request.files["pdf_b"]
    filename = f.filename or "term_work.pdf"
    ok, err = validate_ext(filename, ALL_TERM_WORK_EXTS, "Term Work template")
    if not ok:
        return jsonify({"error": err}), 400
    try:
        raw_bytes = f.read()
        pdf_b_bytes = to_pdf_bytes(raw_bytes, filename)
        titles = json.loads(request.form.get("titles_json", "[]"))
        form = json.loads(request.form.get("form_json", "{}"))
    except Exception as e:
        return jsonify({"error": "Bad payload: " + str(e)}), 400

    if not titles:
        return jsonify({"error": "No practicals selected"}), 400

    try:
        merged = PdfWriter()
        for item in titles:
            fc = dict(form)
            fc["practical_no"] = str(item["sr_no"])
            page_bytes = build_pdf(pdf_b_bytes, item["title"], fc)
            for pg in PdfReader(io.BytesIO(page_bytes)).pages:
                merged.add_page(pg)

        out = io.BytesIO()
        merged.write(out)
        out.seek(0)

        name = make_download_name([t["sr_no"] for t in titles])
        resp = send_file(out, mimetype="application/pdf",
                         as_attachment=True, download_name=name)
        # Lets the front-end read the branded name if it wants to
        resp.headers["X-Download-Name"] = name
        resp.headers["Access-Control-Expose-Headers"] = "Content-Disposition, X-Download-Name"
        return resp
    except Exception as e:
        msg = str(e)
        if "Root object" in msg or "EOF" in msg:
            msg = "Your Term Work file appears corrupted. Please re-upload it."
        return jsonify({"error": msg}), 500


if __name__ == "__main__":
    app.run(debug=True, port=5000)
